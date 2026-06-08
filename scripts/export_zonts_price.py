from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook


ROOT_DIR = Path(__file__).resolve().parent.parent
SOURCE_PATH = Path("/home/drumbler/Документы/Рабочие вопросики/Прайс-лист Финист 12.05.2026.xlsm")
TARGET_PATH = ROOT_DIR / "Товары из прайса в базу.xlsx"
SHEET_NAME = "Зонты"

VALID_PRODUCTS = {
    "ЗПВН-01",
    "ЗПВН-02",
    "ЗПВН-03",
    "ЗПВН-04/01",
    "ЗПВН-04/02",
    "ЗПВН-04/03",
    "ЗПВН-04/05",
    "ЗПВН-05",
    "ЗВН-01",
    "ЗВН-02",
    "ЗВН-03",
    "ЗВН-04/01",
    "ЗВН-04/02",
    "ЗВН-04/03",
    "ЗВН-04/05",
    "ЗВН-05",
    "ЗПВН-01 Премиум",
    "ЗПВН-02 Премиум",
    "ЗПВН-03 Премиум",
    "ЗПВН-04/01 Премиум",
    "ЗПВН-04/02 Премиум",
    "ЗПВН-04/03 Премиум",
    "ЗПВН-04/05 Премиум",
    "ЗПВН-05 Премиум",
    "ЗВН-01 Премиум",
    "ЗВН-02 Премиум",
    "ЗВН-03 Премиум",
    "ЗВН-04/01 Премиум",
    "ЗВН-04/02 Премиум",
    "ЗВН-04/03 Премиум",
    "ЗВН-04/05 Премиум",
    "ЗВН-05 Премиум",
}

HEADER_NORMALIZATION = {
    "ЗПВН-04 Премиум": "ЗПВН-04/03 Премиум",
    "ЗВН-04 Премиум": "ЗВН-04/03 Премиум",
    "ЗВПН-04/01 Премиум": "ЗПВН-04/01 Премиум",
}

SIZE_PATTERN = re.compile(r"^\s*(\d+)\s*[xх*]\s*(\d+)\s*$", re.IGNORECASE)
HEIGHT_PATTERN = re.compile(r"(\d+)")


def normalize_header(value: object) -> str | None:
    if not isinstance(value, str):
        return None

    normalized = " ".join(value.split())
    normalized = HEADER_NORMALIZATION.get(normalized, normalized)
    if normalized in VALID_PRODUCTS:
        return normalized
    return None


def extract_height(value: object) -> int | None:
    if not isinstance(value, str):
        return None
    match = HEIGHT_PATTERN.search(value)
    if match is None:
        return None
    return int(match.group(1))


def extract_size(value: object) -> tuple[int, int] | None:
    if not isinstance(value, str):
        return None
    match = SIZE_PATTERN.match(value)
    if match is None:
        return None
    return int(match.group(1)), int(match.group(2))


def iter_products() -> Iterable[tuple[str, str, float | int]]:
    workbook = load_workbook(SOURCE_PATH, read_only=True, data_only=True, keep_vba=True)
    worksheet = workbook[SHEET_NAME]

    current_product: str | None = None
    current_height: int | None = None
    seen: set[tuple[str, str]] = set()

    for product_cell, height_cell, size_cell, price_cell in worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=4,
        values_only=True,
    ):
        normalized_product = normalize_header(product_cell)
        if normalized_product is not None:
            current_product = normalized_product

        height = extract_height(height_cell)
        if height is not None:
            current_height = height

        if current_product is None or current_height is None:
            continue

        size = extract_size(size_cell)
        if size is None or not isinstance(price_cell, (int, float)):
            continue

        width, depth = size
        characteristic = f"{width}*{depth}*{current_height}"
        dedupe_key = (current_product, characteristic)
        if dedupe_key in seen:
            continue

        seen.add(dedupe_key)
        yield current_product, characteristic, price_cell


def save_to_excel(rows: Iterable[tuple[str, str, float | int]]) -> int:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Зонты"
    worksheet.append(("Изделие", "Характеристика", "Цена"))

    count = 0
    for product, characteristic, price in rows:
        worksheet.append((product, characteristic, price))
        count += 1

    worksheet.column_dimensions["A"].width = 22
    worksheet.column_dimensions["B"].width = 20
    worksheet.column_dimensions["C"].width = 18

    workbook.save(TARGET_PATH)
    return count


def main() -> None:
    count = save_to_excel(iter_products())
    print(f"Created: {TARGET_PATH}")
    print(f"Rows: {count}")


if __name__ == "__main__":
    main()
