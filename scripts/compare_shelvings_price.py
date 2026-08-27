from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import psycopg2
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from app.data.DBrepository import (
    normalize_calculated_product_parameters,
    normalize_calculated_product_series,
    round_calculated_product_cost,
)


DEFAULT_PRICE_PATH = Path(
    "/home/drumbler/Drumbler/Рабочие вопросики/Прайс-лист Финист 12.05.2026.xlsm"
)
SHEET_NAME = "Стеллажи"
DEFAULT_OUTPUT_PATH = Path("shelvings_price_comparison.csv")
DEFAULT_BACKUP_PATH = Path("calculated_products_shelvings_backup.csv")

STANDARD_RAL = "RAL7024"
NO_RAL = "-"
SHELF_TYPE = "С"
ADDITIONAL_REINFORCEMENT = "усиление.нет"
WELD = "разборн"
WELDED_SERIES = {
    "Стеллаж СтПП",
    "Стеллаж СтПЭ",
}
STANDS = "ст.опоры"
PAINTED_PILLAR_SERIES = {
    "Стеллаж СтПБ",
    "Стеллаж СтПЭ",
    "Стеллаж СтПЭр",
}


MATERIAL_ABBR_BY_PRICE_TEXT = {
    "оцинк.ст. 0,7мм": "оц.ст.07",
    "крашеная ч/ст. 0,7мм": "краш.ст.07",
    "нерж.ст. 0,8мм": "н.ст.08",
    "оцинк.ст. 1,0мм": "оц.ст.10",
    "крашеная ч/с 1,0мм": "краш.ст.10",
    "нерж.ст. 1,0мм": "н.ст.10",
}


@dataclass(frozen=True, slots=True)
class ShelvingPriceRow:
    series: str
    parameters: str
    cost: int
    source_row: int
    source_col: str
    price_header: str
    material_text: str
    shelf_material: str

    @property
    def normalized_key(self) -> tuple[str, str]:
        return (
            normalize_calculated_product_series(self.series),
            normalize_calculated_product_parameters(self.parameters),
        )


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _to_int(value: object) -> int:
    return int(float(value))


def _normalize_material_text(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(value.strip().split())


def _shelf_material_from_price_text(value: object) -> str | None:
    return MATERIAL_ABBR_BY_PRICE_TEXT.get(_normalize_material_text(value))


def _series_from_header(header: str) -> str:
    match = re.search(r"СтП[А-Яа-я]+", header)
    if match is None:
        raise ValueError(f"Cannot parse shelving series from header: {header!r}")
    return f"Стеллаж {match.group(0)}"


def _shelf_count_from_header(header: str) -> int:
    match = re.search(r"\((\d+)\s+спл", header)
    if match is None:
        raise ValueError(f"Cannot parse shelf count from header: {header!r}")
    return int(match.group(1))


def _build_standard_parameters(
    *,
    series: str,
    width: int,
    depth: int,
    height: int,
    shelf_material: str,
    shelf_count: int,
) -> str:
    ral_pillars = STANDARD_RAL if series in PAINTED_PILLAR_SERIES else NO_RAL
    ral_shelfs = STANDARD_RAL if shelf_material.startswith("краш.ст") else NO_RAL
    weld = "сварн" if series in WELDED_SERIES else WELD
    return "/".join(
        (
            str(width),
            str(depth),
            str(height),
            ral_pillars,
            shelf_material,
            ral_shelfs,
            str(shelf_count),
            SHELF_TYPE * shelf_count,
            ADDITIONAL_REINFORCEMENT,
            weld,
            STANDS,
        )
    )


def _main_price_columns(rows: list[tuple[object, ...]]) -> list[dict[str, object]]:
    header_row = rows[11]
    height_row = rows[12]
    columns: list[dict[str, object]] = []
    current_header = None

    for col_idx in range(7, 48):
        header_value = header_row[col_idx] if col_idx < len(header_row) else None
        if isinstance(header_value, str) and "СтП" in header_value:
            current_header = header_value

        height = height_row[col_idx] if col_idx < len(height_row) else None
        if current_header and _is_number(height):
            columns.append(
                {
                    "index": col_idx,
                    "letter": _column_letter(col_idx),
                    "series": _series_from_header(current_header),
                    "shelf_count": _shelf_count_from_header(current_header),
                    "height": _to_int(height),
                    "header": current_header,
                }
            )

    return columns


def _column_letter(zero_based_col_idx: int) -> str:
    col = zero_based_col_idx + 1
    result = ""
    while col:
        col, remainder = divmod(col - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _parse_main_table(rows: list[tuple[object, ...]]) -> list[ShelvingPriceRow]:
    price_columns = _main_price_columns(rows)
    parsed_rows: list[ShelvingPriceRow] = []
    for material_text, group in _split_main_width_groups(rows):
        width_values = [_to_int(row[1]) for _, row in group if _is_number(row[1])]
        if not width_values:
            continue
        current_width = width_values[0]

        shelf_material = _shelf_material_from_price_text(material_text)
        if shelf_material is None:
            continue

        for row_number, row in group:
            depth = _to_int(row[2])
            for column in price_columns:
                cost_value = row[column["index"]]
                if not _is_number(cost_value):
                    continue

                series = str(column["series"])
                params = _build_standard_parameters(
                    series=series,
                    width=current_width,
                    depth=depth,
                    height=int(column["height"]),
                    shelf_material=shelf_material,
                    shelf_count=int(column["shelf_count"]),
                )
                parsed_rows.append(
                    ShelvingPriceRow(
                        series=series,
                        parameters=params,
                        cost=round_calculated_product_cost(cost_value),
                        source_row=row_number,
                        source_col=str(column["letter"]),
                        price_header=str(column["header"]),
                        material_text=material_text,
                        shelf_material=shelf_material,
                    )
                )

    return parsed_rows


def _split_main_width_groups(
    rows: list[tuple[object, ...]],
) -> Iterable[tuple[str, list[tuple[int, tuple[object, ...]]]]]:
    group: list[tuple[int, tuple[object, ...]]] = []
    current_material_text = None
    previous_depth = None

    for row_number in range(21, 224):
        row = rows[row_number - 1]
        material_text = _normalize_material_text(row[0])
        if material_text in MATERIAL_ABBR_BY_PRICE_TEXT:
            if group and current_material_text is not None:
                yield current_material_text, group
                group = []
            current_material_text = material_text
            previous_depth = None

        if current_material_text is None or not _is_number(row[2]):
            continue

        depth = _to_int(row[2])
        if group and previous_depth is not None and depth <= previous_depth:
            yield current_material_text, group
            group = []

        group.append((row_number, row))
        previous_depth = depth

    if group and current_material_text is not None:
        yield current_material_text, group


def _split_stpp_width_groups(rows: list[tuple[object, ...]]) -> Iterable[list[tuple[int, tuple[object, ...]]]]:
    group: list[tuple[int, tuple[object, ...]]] = []
    previous_depth = None

    for row_number in range(277, 340):
        row = rows[row_number - 1]
        depth = row[2]
        if not _is_number(depth):
            continue

        if group and previous_depth is not None and _to_int(depth) <= previous_depth:
            yield group
            group = []

        group.append((row_number, row))
        previous_depth = _to_int(depth)

    if group:
        yield group


def _parse_stpp_table(rows: list[tuple[object, ...]]) -> list[ShelvingPriceRow]:
    series = "Стеллаж СтПП"
    shelf_count = 4
    header = str(rows[273][3])
    height_columns = [
        {"index": col_idx, "letter": _column_letter(col_idx), "height": _to_int(rows[275][col_idx])}
        for col_idx in range(3, 6)
        if _is_number(rows[275][col_idx])
    ]
    parsed_rows: list[ShelvingPriceRow] = []
    current_material_text = None

    for group in _split_stpp_width_groups(rows):
        width_values = [_to_int(row[1]) for _, row in group if _is_number(row[1])]
        if not width_values:
            continue
        width = width_values[0]

        for row_number, row in group:
            material_text = _normalize_material_text(row[0])
            if material_text in MATERIAL_ABBR_BY_PRICE_TEXT:
                current_material_text = material_text

            shelf_material = _shelf_material_from_price_text(current_material_text)
            if shelf_material is None:
                continue

            depth = _to_int(row[2])
            for column in height_columns:
                cost_value = row[column["index"]]
                if not _is_number(cost_value):
                    continue

                params = _build_standard_parameters(
                    series=series,
                    width=width,
                    depth=depth,
                    height=int(column["height"]),
                    shelf_material=shelf_material,
                    shelf_count=shelf_count,
                )
                parsed_rows.append(
                    ShelvingPriceRow(
                        series=series,
                        parameters=params,
                        cost=round_calculated_product_cost(cost_value),
                        source_row=row_number,
                        source_col=str(column["letter"]),
                        price_header=header,
                        material_text=str(current_material_text),
                        shelf_material=shelf_material,
                    )
                )

    return parsed_rows


def parse_shelving_price_rows(price_path: Path) -> list[ShelvingPriceRow]:
    workbook = load_workbook(price_path, read_only=True, data_only=True, keep_vba=False)
    worksheet = workbook[SHEET_NAME]
    rows = list(worksheet.iter_rows(min_row=1, max_row=340, max_col=50, values_only=True))
    workbook.close()
    return _parse_main_table(rows) + _parse_stpp_table(rows)


def fetch_existing_calculated_products() -> dict[tuple[str, str], int]:
    load_dotenv(dotenv_path=".env")
    conn_data = {
        "dbname": os.environ.get("DB_NAME"),
        "user": os.environ.get("DB_USER"),
        "password": os.environ.get("DB_PASS"),
        "host": os.environ.get("DB_IP"),
        "port": os.environ.get("DB_PORT"),
    }
    with psycopg2.connect(**conn_data) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                    SELECT series, parameters, cost
                    FROM calculated_products
                    WHERE UPPER(series) LIKE 'СТЕЛЛАЖ%'
                """
            )
            return {
                (
                    normalize_calculated_product_series(series),
                    normalize_calculated_product_parameters(parameters),
                ): round_calculated_product_cost(cost)
                for series, parameters, cost in cur.fetchall()
            }


def backup_existing_calculated_products(output_path: Path) -> int:
    load_dotenv(dotenv_path=".env")
    conn_data = {
        "dbname": os.environ.get("DB_NAME"),
        "user": os.environ.get("DB_USER"),
        "password": os.environ.get("DB_PASS"),
        "host": os.environ.get("DB_IP"),
        "port": os.environ.get("DB_PORT"),
    }
    with psycopg2.connect(**conn_data) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                    SELECT series, parameters, cost
                    FROM calculated_products
                    WHERE UPPER(series) LIKE 'СТЕЛЛАЖ%'
                    ORDER BY series, parameters
                """
            )
            rows = cur.fetchall()

    with output_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(("series", "parameters", "cost"))
        writer.writerows(rows)

    return len(rows)


def insert_missing_price_rows(
    price_rows: list[ShelvingPriceRow],
    existing_rows: dict[tuple[str, str], int],
) -> int:
    rows_to_insert = []
    seen = set()
    for row in price_rows:
        if row.normalized_key in existing_rows or row.normalized_key in seen:
            continue
        seen.add(row.normalized_key)
        rows_to_insert.append((row.normalized_key[0], row.normalized_key[1], row.cost))

    if not rows_to_insert:
        return 0

    load_dotenv(dotenv_path=".env")
    conn_data = {
        "dbname": os.environ.get("DB_NAME"),
        "user": os.environ.get("DB_USER"),
        "password": os.environ.get("DB_PASS"),
        "host": os.environ.get("DB_IP"),
        "port": os.environ.get("DB_PORT"),
    }
    with psycopg2.connect(**conn_data) as conn:
        with conn.cursor() as cur:
            cur.executemany(
                """
                    INSERT INTO calculated_products (series, parameters, cost)
                    VALUES (%s, %s, %s)
                    ON CONFLICT (series, parameters) DO NOTHING
                """,
                rows_to_insert,
            )

    return len(rows_to_insert)


def update_different_price_rows(
    price_rows: list[ShelvingPriceRow],
    existing_rows: dict[tuple[str, str], int],
) -> int:
    rows_to_update = []
    seen = set()
    for row in price_rows:
        if row.normalized_key in seen:
            continue
        seen.add(row.normalized_key)

        db_cost = existing_rows.get(row.normalized_key)
        if db_cost is None or db_cost == row.cost:
            continue

        rows_to_update.append((row.cost, row.normalized_key[0], row.normalized_key[1]))

    if not rows_to_update:
        return 0

    load_dotenv(dotenv_path=".env")
    conn_data = {
        "dbname": os.environ.get("DB_NAME"),
        "user": os.environ.get("DB_USER"),
        "password": os.environ.get("DB_PASS"),
        "host": os.environ.get("DB_IP"),
        "port": os.environ.get("DB_PORT"),
    }
    with psycopg2.connect(**conn_data) as conn:
        with conn.cursor() as cur:
            cur.executemany(
                """
                    UPDATE calculated_products
                    SET cost = %s
                    WHERE series = %s AND parameters = %s
                """,
                rows_to_update,
            )

    return len(rows_to_update)


def write_comparison(
    price_rows: list[ShelvingPriceRow],
    existing_rows: dict[tuple[str, str], int],
    output_path: Path,
) -> dict[str, int]:
    seen: dict[tuple[str, str], ShelvingPriceRow] = {}
    duplicate_keys: set[tuple[str, str]] = set()
    for row in price_rows:
        if row.normalized_key in seen:
            duplicate_keys.add(row.normalized_key)
        seen[row.normalized_key] = row

    summary = {
        "price_rows": len(price_rows),
        "unique_price_rows": len(seen),
        "duplicates_in_price": len(duplicate_keys),
        "same": 0,
        "missing_in_db": 0,
        "different_cost": 0,
        "db_only": 0,
    }

    fieldnames = [
        "status",
        "series",
        "parameters",
        "price_cost",
        "db_cost",
        "diff",
        "source_row",
        "source_col",
        "price_header",
        "material_text",
        "shelf_material",
    ]

    with output_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()

        for key, row in sorted(seen.items()):
            db_cost = existing_rows.get(key)
            if key in duplicate_keys:
                status = "duplicate_in_price"
            elif db_cost is None:
                status = "missing_in_db"
                summary["missing_in_db"] += 1
            elif db_cost == row.cost:
                status = "same"
                summary["same"] += 1
            else:
                status = "different_cost"
                summary["different_cost"] += 1

            writer.writerow(
                {
                    "status": status,
                    "series": row.normalized_key[0],
                    "parameters": row.normalized_key[1],
                    "price_cost": row.cost,
                    "db_cost": "" if db_cost is None else db_cost,
                    "diff": "" if db_cost is None else row.cost - db_cost,
                    "source_row": row.source_row,
                    "source_col": row.source_col,
                    "price_header": row.price_header,
                    "material_text": row.material_text,
                    "shelf_material": row.shelf_material,
                }
            )

        for key, db_cost in sorted(existing_rows.items()):
            if key in seen:
                continue
            summary["db_only"] += 1
            writer.writerow(
                {
                    "status": "db_only",
                    "series": key[0],
                    "parameters": key[1],
                    "price_cost": "",
                    "db_cost": db_cost,
                    "diff": "",
                    "source_row": "",
                    "source_col": "",
                    "price_header": "",
                    "material_text": "",
                    "shelf_material": "",
                }
            )

    return summary


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Dry-run comparison for standard shelving prices."
    )
    parser.add_argument("--price-path", type=Path, default=DEFAULT_PRICE_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    parser.add_argument("--insert-missing", action="store_true")
    parser.add_argument("--update-different", action="store_true")
    parser.add_argument("--backup", type=Path, default=DEFAULT_BACKUP_PATH)
    args = parser.parse_args()

    price_rows = parse_shelving_price_rows(args.price_path)
    existing_rows = fetch_existing_calculated_products()
    if args.insert_missing or args.update_different:
        backup_count = backup_existing_calculated_products(args.backup)
        print(f"backup: {args.backup}")
        print(f"backup_rows: {backup_count}")

    if args.insert_missing:
        inserted_count = insert_missing_price_rows(price_rows, existing_rows)
        print(f"inserted_missing_rows: {inserted_count}")
        existing_rows = fetch_existing_calculated_products()

    if args.update_different:
        updated_count = update_different_price_rows(price_rows, existing_rows)
        print(f"updated_different_rows: {updated_count}")
        existing_rows = fetch_existing_calculated_products()

    summary = write_comparison(price_rows, existing_rows, args.output)

    print(f"price_path: {args.price_path}")
    print(f"output: {args.output}")
    for key, value in summary.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
