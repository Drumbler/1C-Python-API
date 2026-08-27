from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


ROOT_DIR = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = ROOT_DIR / "Выгрузка цен — ванны и рукомойники 12.05.2026.xlsx"
LOAD_WORKBOOK_PATH = ROOT_DIR / "Ванны со стандартными характеристиками для загрузки.xlsx"
SOURCE_SHEET = "Выгрузка"
IMPORT_SHEET = "Ванны — для загрузки"
CHARACTERISTIC_HEADER = "Стандартная характеристика"

SERIES_PATTERN = re.compile(r"^В[МК]([БЛНПСУЭ]).*-\d$", re.IGNORECASE)


def normalize_compartment(value: object) -> str | None:
    if not isinstance(value, str) or not value.strip():
        return None

    dimensions = re.findall(r"\d+", value)
    if len(dimensions) != 3:
        raise ValueError(f"Некорректный размер моечного отделения: {value!r}")
    return "х".join(dimensions)


def define_assembly(series: str, welded_marker: object = None) -> str:
    normalized = series.strip().upper()
    match = SERIES_PATTERN.match(normalized)
    if match is None:
        raise ValueError(f"Неподдерживаемая серия ванны: {series!r}")

    series_letter = match.group(1)
    if series_letter in {"П", "С", "Э"}:
        return "сварн"
    if series_letter in {"Б", "Л", "Н", "У"}:
        return "сборн"
    return "сварн" if str(welded_marker).strip().casefold() == "да" else "сборн"


def define_standard_color(series: str) -> str:
    match = SERIES_PATTERN.match(series.strip())
    if match is None:
        raise ValueError(f"Неподдерживаемая серия ванны: {series!r}")
    return "RAL7024" if match.group(1).upper() == "Э" else "-"


def build_characteristic(
    series: str,
    width: object,
    depth: object,
    height: object,
    compartment: object = None,
    welded_marker: object = None,
) -> str:
    dimensions = [str(value).strip() for value in (width, depth, height)]
    if any(not value or not value.isdigit() for value in dimensions):
        raise ValueError(
            f"Некорректные габариты для {series}: {width!r}/{depth!r}/{height!r}"
        )

    parts = [*dimensions, "глуб.м.о(300мм)", define_standard_color(series)]
    normalized_compartment = normalize_compartment(compartment)
    if normalized_compartment is not None:
        parts.append(f"моеч.отд.{normalized_compartment}")
    parts.extend(
        (
            "борт.НЕТ",
            "смес.НЕТ",
            "-",
            "ст.опоры",
            "слив.ст",
            "полка.НЕТ",
            define_assembly(series, welded_marker),
        )
    )
    return "/".join(parts)


def save_load_workbook(
    rows: list[tuple[str, str, float | int]],
    target_path: Path = LOAD_WORKBOOK_PATH,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ванны"
    worksheet.append(("Изделие", "Характеристика", "Цена"))
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        worksheet.append(row)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 105
    worksheet.column_dimensions["C"].width = 18
    workbook.save(target_path)
    workbook.close()


def add_characteristics(workbook_path: Path) -> int:
    workbook = load_workbook(workbook_path)
    worksheet = workbook[SOURCE_SHEET]

    header_values = [cell.value for cell in worksheet[1]]
    if CHARACTERISTIC_HEADER in header_values:
        characteristic_column = header_values.index(CHARACTERISTIC_HEADER) + 1
    else:
        characteristic_column = worksheet.max_column + 1
    worksheet.cell(1, characteristic_column, CHARACTERISTIC_HEADER)
    worksheet.cell(1, characteristic_column).font = Font(bold=True)
    worksheet.column_dimensions[
        worksheet.cell(1, characteristic_column).column_letter
    ].width = 105
    for row_number in range(2, worksheet.max_row + 1):
        worksheet.cell(row_number, characteristic_column).value = None

    if IMPORT_SHEET in workbook.sheetnames:
        del workbook[IMPORT_SHEET]
    import_sheet = workbook.create_sheet(IMPORT_SHEET)
    import_sheet.append(("Изделие", "Характеристика", "Цена"))
    for cell in import_sheet[1]:
        cell.font = Font(bold=True)

    load_rows: list[tuple[str, str, float | int]] = []
    for row_number in range(2, worksheet.max_row + 1):
        series = worksheet.cell(row_number, 1).value
        if not isinstance(series, str) or SERIES_PATTERN.match(series.strip()) is None:
            continue

        try:
            characteristic = build_characteristic(
                series=series,
                width=worksheet.cell(row_number, 2).value,
                depth=worksheet.cell(row_number, 3).value,
                height=worksheet.cell(row_number, 4).value,
                compartment=worksheet.cell(row_number, 5).value,
                welded_marker=worksheet.cell(row_number, 7).value,
            )
        except ValueError as error:
            print(f"Пропущена строка {row_number}: {error}")
            continue
        worksheet.cell(row_number, characteristic_column, characteristic)
        import_sheet.append(
            (series, characteristic, worksheet.cell(row_number, 6).value)
        )
        load_rows.append(
            (series, characteristic, worksheet.cell(row_number, 6).value)
        )

    import_sheet.freeze_panes = "A2"
    import_sheet.auto_filter.ref = import_sheet.dimensions
    import_sheet.column_dimensions["A"].width = 18
    import_sheet.column_dimensions["B"].width = 105
    import_sheet.column_dimensions["C"].width = 18

    workbook.save(workbook_path)
    workbook.close()
    save_load_workbook(load_rows)
    return len(load_rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Добавить стандартные характеристики к выгрузке цен ванн."
    )
    parser.add_argument("workbook", nargs="?", type=Path, default=WORKBOOK_PATH)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    count = add_characteristics(args.workbook)
    print(f"Обновлён файл: {args.workbook}")
    print(f"Создан файл для загрузки: {LOAD_WORKBOOK_PATH}")
    print(f"Характеристик: {count}")


if __name__ == "__main__":
    main()
